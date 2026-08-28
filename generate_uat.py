#!/usr/bin/env python3
"""
BestoClinic — UAT Test Case Generator
Generates a comprehensive .xlsx workbook with 4 role-based sheets:
  1. Patient
  2. Doctor & Staff
  3. Clinic Admin
  4. Platform Admin
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date

# ─── Styling ───────────────────────────────────────────────
HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
MODULE_FONT = Font(name="Calibri", bold=True, size=11, color="1F4E79")
MODULE_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
BODY_FONT = Font(name="Calibri", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)

# Priority fill colors
P1_FILL = PatternFill(start_color="FDDEDE", end_color="FDDEDE", fill_type="solid")  # red tint
P2_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")  # yellow tint
P3_FILL = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")  # green tint

STATUS_OPTIONS = '"Not Tested,Pass,Fail,Blocked,Deferred"'

COLUMNS = [
    ("Test Case ID", 14),
    ("Module", 18),
    ("Sub-Module", 18),
    ("Test Case", 48),
    ("Pre-Conditions", 36),
    ("Test Steps", 52),
    ("Expected Result", 40),
    ("Priority", 10),
    ("Status", 12),
    ("Tested By", 14),
    ("Test Date", 12),
    ("Defect ID", 12),
    ("Comments", 30),
]


def setup_sheet(ws, title, columns):
    ws.sheet_properties.tabColor = {
        "Patient": "2E86C1",
        "Doctor & Staff": "27AE60",
        "Clinic Admin": "E67E22",
        "Platform Admin": "8E44AD",
    }.get(title, "333333")

    # Header row
    for col_idx, (col_name, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"


def add_row(ws, row_num, data, is_module_header=False):
    for col_idx, value in enumerate(data, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=value)
        cell.font = MODULE_FONT if is_module_header else BODY_FONT
        cell.alignment = CENTER if col_idx in (1, 8, 9, 10, 11, 12) else WRAP
        cell.border = THIN_BORDER
        if is_module_header:
            cell.fill = MODULE_FILL

    # Priority-based fill on priority column
    priority = data[7] if len(data) > 7 else ""
    if priority == "P1":
        ws.cell(row=row_num, column=8).fill = P1_FILL
    elif priority == "P2":
        ws.cell(row=row_num, column=8).fill = P2_FILL
    elif priority == "P3":
        ws.cell(row=row_num, column=8).fill = P3_FILL

    # Status dropdown
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1=STATUS_OPTIONS, allow_blank=True)
    dv.error = "Please select a valid status"
    dv.errorTitle = "Invalid Status"
    ws.add_data_validation(dv)
    dv.add(ws.cell(row=row_num, column=9))


def add_tests(ws, test_id_prefix, tests):
    """Add a list of test case tuples to the sheet. Returns next row number."""
    row = ws.max_row + 1
    tc_num = 1
    for t in tests:
        if len(t) == 1:
            # Module header row
            add_row(ws, row, [t[0]] + [""] * 12, is_module_header=True)
        else:
            tc_id = f"{test_id_prefix}-{tc_num:03d}"
            # t = (module, submodule, test_case, preconditions, steps, expected, priority)
            add_row(ws, row, [tc_id] + list(t) + ["Not Tested", "", "", "", ""])
            tc_num += 1
        row += 1
    return row


# ═══════════════════════════════════════════════════════════
#  TEST CASE DATA
# ═══════════════════════════════════════════════════════════

PATIENT_TESTS = [
    # ── Auth ──
    ("AUTH",),
    ("Auth", "Registration", "Register new patient account with valid details",
     "App installed, clinic selected", "1. Open app\n2. Tap 'Register'\n3. Fill name, email, phone, password\n4. Submit",
     "Account created, redirected to dashboard", "P1"),
    ("Auth", "Registration", "Register with duplicate email shows error",
     "Email already registered", "1. Open Register\n2. Enter existing email\n3. Submit",
     "Error: 'Email already registered'", "P1"),
    ("Auth", "Registration", "Register with weak password shows validation error",
     "App installed", "1. Open Register\n2. Enter password < 8 chars\n3. Submit",
     "Validation error displayed", "P2"),
    ("Auth", "Login", "Login with valid patient credentials",
     "Registered patient account", "1. Open Login\n2. Enter email & password\n3. Tap Login",
     "Logged in, dashboard loads with correct data", "P1"),
    ("Auth", "Login", "Login with invalid credentials shows error",
     "App on login screen", "1. Enter wrong email/password\n2. Tap Login",
     "Error: 'Invalid credentials'", "P1"),
    ("Auth", "Login", "Login with suspended tenant shows subscription blocked message",
     "Tenant suspended by platform admin", "1. Enter valid credentials\n2. Tap Login",
     "Error: Subscription blocked with reason", "P2"),
    ("Auth", "Clinic Selection", "Select clinic from available tenants list",
     "Multiple tenants configured", "1. Open app\n2. See clinic list\n3. Tap a clinic",
     "Clinic selected, tenant slug stored, proceed to login", "P1"),
    ("Auth", "Forgot Password", "Request password reset via email",
     "Registered account", "1. Tap 'Forgot Password'\n2. Enter email\n3. Submit",
     "Success message: reset email sent", "P2"),
    ("Auth", "Logout", "Logout clears session and navigates to auth",
     "Logged in as patient", "1. Go to Profile\n2. Tap Logout\n3. Confirm",
     "Session cleared, socket disconnected, back to login screen", "P1"),
    ("Auth", "Token Refresh", "App refreshes expired access token automatically",
     "Token about to expire", "1. Wait for token expiry\n2. Make any API call",
     "Token refreshed silently, API call succeeds", "P2"),

    # ── Dashboard ──
    ("DASHBOARD",),
    ("Dashboard", "Stats", "Dashboard displays correct appointment, prescription, payment, and notification counts",
     "Logged in, has data", "1. Navigate to Dashboard\n2. Check all 4 stat cards",
     "Counts match backend data (upcoming appointments, active prescriptions, pending payments, unread notifications)", "P1"),
    ("Dashboard", "Stats", "Stat cards navigate to respective screens on tap",
     "Dashboard loaded", "1. Tap 'Upcoming Appointments' card\n2. Tap back\n3. Tap each card",
     "Each card navigates to correct screen (Appointments, Records/Prescriptions, Payments, Notifications)", "P2"),
    ("Dashboard", "Next Appointment", "Next appointment card shows correct info",
     "Has upcoming appointment", "1. View Dashboard\n2. Check next appointment card",
     "Shows doctor name, specialization, type, date/time", "P1"),
    ("Dashboard", "Next Appointment", "Tapping next appointment navigates to detail",
     "Next appointment visible", "1. Tap the next appointment card",
     "Navigates to AppointmentDetail screen with correct data", "P2"),
    ("Dashboard", "Quick Actions", "Book Appointment quick action works",
     "Dashboard loaded", "1. Tap 'Book Appointment' quick action",
     "Navigates to booking wizard", "P1"),
    ("Dashboard", "Quick Actions", "Emergency quick action initiates emergency chat",
     "Dashboard loaded", "1. Tap 'Emergency' quick action",
     "Creates emergency session, opens EmergencyChat modal", "P1"),
    ("Dashboard", "Quick Actions", "Share Records quick action works",
     "Dashboard loaded", "1. Tap 'Share Records' quick action",
     "Navigates to ShareRecords screen", "P2"),
    ("Dashboard", "Refresh", "Pull-to-refresh updates dashboard data",
     "Dashboard loaded", "1. Pull down to refresh",
     "Dashboard data reloads with fresh counts", "P2"),
    ("Dashboard", "Real-time", "Notification count updates in real-time via WebSocket",
     "Dashboard open, socket connected", "1. Trigger notification from backend\n2. Observe dashboard",
     "Notification count increments without manual refresh", "P2"),
    ("Dashboard", "Guardian", "Guardian can switch between patient profiles",
     "Logged in as guardian with multiple profiles", "1. Tap profile switcher\n2. Select different profile",
     "Dashboard reloads with selected patient's data", "P1"),

    # ── Appointments ──
    ("APPOINTMENTS",),
    ("Appointments", "List", "View all appointments with pagination",
     "Has appointments", "1. Go to Appointments tab\n2. Scroll through list",
     "Appointments listed with doctor, date, status, type; paginated", "P1"),
    ("Appointments", "Filters", "Filter appointments by Upcoming/Completed/Cancelled",
     "Has mixed-status appointments", "1. Tap 'Upcoming' chip\n2. Tap 'Completed'\n3. Tap 'Cancelled'",
     "List filters correctly; Upcoming = requested+confirmed+checked_in+in_progress", "P1"),
    ("Appointments", "Detail", "View full appointment detail",
     "Has an appointment", "1. Tap an appointment card",
     "Shows status, doctor info, schedule (date/time/duration/branch), reason, notes", "P1"),
    ("Appointments", "Book", "Book new appointment via 4-step wizard",
     "Logged in as patient", "1. Tap Book Appointment\n2. Select Department\n3. Select Doctor\n4. Choose date/time/type\n5. Confirm",
     "Appointment created with 'requested' status, confirmation shown", "P1"),
    ("Appointments", "Book", "Booking wizard validates required fields at each step",
     "On booking wizard", "1. Try proceeding without selecting department\n2. Try without doctor\n3. Try without date/time",
     "Validation errors shown at each step", "P2"),
    ("Appointments", "Book", "Book teleconsult appointment type",
     "Department selected, doctor selected", "1. Select 'Teleconsult' type\n2. Pick date/time\n3. Confirm",
     "Teleconsult appointment created, 'Join Video Call' button appears when confirmed", "P1"),
    ("Appointments", "Cancel", "Cancel a requested/confirmed appointment with reason",
     "Has cancellable appointment (requested/confirmed)", "1. Open appointment detail\n2. Tap Cancel\n3. Enter reason\n4. Confirm",
     "Status changes to 'cancelled', reason saved, list updates", "P1"),
    ("Appointments", "Cancel", "Cancel button not shown for completed/cancelled appointments",
     "Has completed appointment", "1. Open completed appointment detail",
     "Cancel button is NOT visible", "P2"),
    ("Appointments", "Reschedule", "Reschedule a requested/confirmed appointment",
     "Has reschedulable appointment", "1. Open detail\n2. Tap Reschedule\n3. Pick new date\n4. Pick time slot\n5. Confirm",
     "Appointment rescheduled, new date/time shown", "P1"),
    ("Appointments", "Reschedule", "Reschedule option hidden for non-reschedulable statuses",
     "Has checked_in appointment", "1. Open appointment detail",
     "Reschedule button is NOT visible", "P2"),

    # ── Video Call (Telemedicine) ──
    ("VIDEO CALL / TELEMEDICINE",),
    ("Video Call", "Join", "Join teleconsult appointment as patient",
     "Teleconsult appointment in confirmed/checked_in/in_progress status", "1. Open teleconsult appointment detail\n2. Tap 'Join Video Call'\n3. Grant camera/mic permissions",
     "'Join Video Call' button visible; video call screen opens, local camera preview shown, waiting for doctor", "P1"),
    ("Video Call", "Permissions", "Camera/mic permission request on first join",
     "Permissions not yet granted", "1. Tap 'Join Video Call'\n2. System permission dialog appears",
     "Android: CAMERA + RECORD_AUDIO requested; iOS: system prompts shown", "P1"),
    ("Video Call", "Permissions", "Permission denial shows error and settings link",
     "Denied camera/mic permissions", "1. Deny permissions on prompt",
     "Error message displayed: 'Camera and microphone permissions are required'", "P2"),
    ("Video Call", "Connect", "P2P video connection established when both parties join",
     "Patient and doctor both joined same session", "1. Patient joins video call\n2. Doctor joins video call",
     "WebRTC connection established, both see each other's video, audio works", "P1"),
    ("Video Call", "Controls", "Toggle mute/unmute during call",
     "In active video call", "1. Tap microphone button to mute\n2. Tap again to unmute",
     "Audio toggles, remote party sees mute indicator, icon changes", "P1"),
    ("Video Call", "Controls", "Toggle camera on/off during call",
     "In active video call", "1. Tap camera button to turn off\n2. Tap again to turn on",
     "Video track toggles, local PIP hides when off, remote sees black frame", "P1"),
    ("Video Call", "Controls", "Switch front/back camera",
     "In active video call with camera on", "1. Tap camera flip button",
     "Camera switches between front and rear", "P2"),
    ("Video Call", "Controls", "End call button terminates session",
     "In active video call", "1. Tap end call (red hangup) button",
     "Call ends, 'Call Ended' shown with duration, navigates back after 2s, session status → ended", "P1"),
    ("Video Call", "Duration", "Call duration timer displays correctly",
     "Connected video call", "1. Observe duration timer after connection",
     "Timer starts at 00:00 and increments every second", "P2"),
    ("Video Call", "Duration", "Auto-end when max duration reached",
     "Call approaching max duration", "1. Let call run to maxDuration limit",
     "Call auto-ends when duration >= maxDuration", "P3"),
    ("Video Call", "Reconnect", "Reconnection banner shown on ICE disconnect",
     "In active call, network briefly drops", "1. Briefly disable network\n2. Re-enable",
     "Yellow banner: 'Connection lost — reconnecting...' shown, reconnects automatically", "P2"),
    ("Video Call", "Remote", "Remote participant mute indicator shown",
     "Connected call, remote party mutes", "1. Remote party mutes microphone",
     "Badge appears: '{Name} muted' with mic-off icon", "P3"),
    ("Video Call", "Cleanup", "Navigating away cleans up WebRTC resources",
     "In video call", "1. Press hardware back or end call",
     "Local stream stopped, peer connection closed, signaling room left, store reset", "P2"),

    # ── Emergency ──
    ("EMERGENCY",),
    ("Emergency", "Create", "Create emergency chat session from dashboard",
     "Logged in as patient", "1. Dashboard → Emergency quick action\n2. Confirm",
     "Emergency session created, EmergencyChat modal opens", "P1"),
    ("Emergency", "Chat", "Send and receive messages in emergency chat",
     "Emergency chat open, doctor accepted", "1. Type message\n2. Tap send\n3. Wait for doctor reply",
     "Message appears in chat (right-aligned), doctor reply appears (left-aligned) in real-time", "P1"),
    ("Emergency", "Chat", "Emergency video type redirects to VideoCall",
     "Emergency created with type='video'", "1. Emergency of type video starts",
     "Auto-redirects to VideoCall screen", "P2"),

    # ── Prescriptions ──
    ("PRESCRIPTIONS & MEDICAL RECORDS",),
    ("Records", "Prescriptions", "View list of prescriptions with status",
     "Has prescriptions", "1. Go to Records tab → Prescriptions tab",
     "Prescriptions listed with medication names, doctor, date, status (active/cancelled)", "P1"),
    ("Records", "Prescriptions", "View prescription detail with medications",
     "Has a prescription", "1. Tap a prescription",
     "Shows diagnosis, all medications (name, dosage, frequency, duration, instructions), doctor, date", "P1"),
    ("Records", "Medical History", "View treatment history timeline",
     "Has treatment records", "1. Go to Records tab → History tab",
     "Timeline shows treatment records with diagnosis, doctor, date in chronological order", "P1"),

    # ── Payments & Billing ──
    ("PAYMENTS & BILLING",),
    ("Payments", "List", "View invoices with filter chips (All/Pending/Paid/Overdue)",
     "Has invoices in various statuses", "1. Go to Payments tab\n2. Tap each filter chip",
     "Invoice list filters correctly by status; shows invoice#, amount (INR), date, status", "P1"),
    ("Payments", "Pay", "Initiate payment for pending invoice",
     "Has pending invoice", "1. Open pending invoice\n2. Tap 'Pay' button",
     "Payment gateway flow initiates (Razorpay order created)", "P1"),
    ("Payments", "Summary", "View financial summary",
     "Has payment history", "1. Navigate to Financial Summary",
     "Shows total outstanding, total paid, payment breakdown", "P2"),
    ("Payments", "PDF", "Download invoice as PDF",
     "Has an invoice", "1. Open invoice detail\n2. Tap download/PDF icon",
     "PDF downloads or opens for viewing", "P3"),

    # ── Notifications ──
    ("NOTIFICATIONS",),
    ("Notifications", "List", "View notification list with types",
     "Has notifications", "1. Go to Notifications screen",
     "Shows notifications with type icon (appointment/payment/system/alert/reminder), title, message, time ago", "P1"),
    ("Notifications", "Read", "Mark single notification as read",
     "Has unread notification", "1. Tap an unread notification",
     "Notification marked as read, visual indicator changes", "P2"),
    ("Notifications", "Read All", "Mark all notifications as read",
     "Has multiple unread notifications", "1. Tap 'Mark All Read' button",
     "All notifications marked read, unread count → 0", "P2"),
    ("Notifications", "Real-time", "New notification appears in real-time",
     "Notification screen open, socket connected", "1. Trigger notification from backend",
     "New notification appears at top of list without manual refresh", "P2"),

    # ── Profile & Settings ──
    ("PROFILE & SETTINGS",),
    ("Profile", "View", "View patient profile with all details",
     "Logged in", "1. Go to Profile tab",
     "Shows avatar, name, patient ID, contact info (email, phone, DOB, address)", "P1"),
    ("Profile", "Edit", "Edit profile details (name, phone, address)",
     "On Profile screen", "1. Tap edit button\n2. Update name/phone/address\n3. Save",
     "Profile updated, changes reflected immediately", "P1"),
    ("Profile", "Consents", "View and manage data sharing consents",
     "Has consents", "1. Profile → Consents\n2. View list",
     "Shows consent type, status, granted/expires dates; approve/revoke buttons work", "P2"),
    ("Profile", "Access Log", "View who accessed my records",
     "Has access log entries", "1. Profile → Access Log (via Medical Records)",
     "Shows accessor name, access type, timestamp, resource accessed", "P2"),
    ("Profile", "Share Records", "Create share token for records",
     "On Share Records screen", "1. Select record type\n2. Set expiry\n3. Create share link",
     "Share token generated with unique link and expiry", "P2"),
    ("Profile", "Review", "Submit review for completed appointment",
     "Has completed appointment", "1. Open completed appointment\n2. Tap 'Write Review'\n3. Select 1-5 stars\n4. Add comment\n5. Submit",
     "Review submitted, confirmation shown", "P2"),
    ("Settings", "Theme", "Switch between Light/Dark/System theme",
     "On Settings screen", "1. Tap Light/Dark/System chip",
     "App theme changes immediately", "P3"),
    ("Settings", "Logout", "Logout from settings screen",
     "On Settings screen", "1. Tap Logout\n2. Confirm",
     "Session cleared, navigated to login", "P2"),
]


DOCTOR_TESTS = [
    # ── Auth ──
    ("AUTH",),
    ("Auth", "Login", "Doctor login with valid staff credentials",
     "Doctor account exists in tenant", "1. Open app\n2. Select clinic\n3. Switch to 'Staff Login'\n4. Enter email/password\n5. Login",
     "Logged in as doctor, doctor dashboard loads", "P1"),
    ("Auth", "Login", "Staff login with non-doctor role (receptionist, billing officer, etc.)",
     "Staff account with receptionist role", "1. Login with staff credentials",
     "Doctor/staff dashboard loads with role-appropriate features", "P1"),
    ("Auth", "Login", "Login rejected for patient credentials on doctor login",
     "Patient account only", "1. Enter patient credentials on doctor login",
     "Error or role mismatch handled gracefully", "P2"),
    ("Auth", "Logout", "Doctor logout clears session",
     "Logged in as doctor", "1. Profile → Logout",
     "Session cleared, socket disconnected, back to login", "P1"),

    # ── Dashboard ──
    ("DASHBOARD",),
    ("Dashboard", "Stats", "Dashboard shows today's appointment count, completed, patients this month, revenue",
     "Logged in as doctor", "1. View dashboard",
     "4 stat cards show correct numbers matching backend /doctor/stats", "P1"),
    ("Dashboard", "Next Appointment", "Next appointment card shows correct patient info",
     "Has upcoming appointment", "1. View dashboard",
     "Shows patient name, appointment type, date/time", "P1"),
    ("Dashboard", "Emergency", "Emergency alert section shows pending emergencies",
     "Pending emergency exists", "1. View dashboard",
     "Emergency alert card(s) shown with patient name, reason, accept/decline buttons", "P1"),
    ("Dashboard", "Emergency", "Accept emergency from dashboard",
     "Pending emergency shown", "1. Tap 'Accept' on emergency alert",
     "Emergency session accepted, navigates to EmergencyChat or VideoCall", "P1"),
    ("Dashboard", "Follow-ups", "Upcoming follow-ups section shows next 3",
     "Has follow-up appointments", "1. View dashboard → follow-ups section",
     "Up to 3 follow-up appointments shown with patient, date; 'See All' navigates to FollowUpList", "P2"),
    ("Dashboard", "Schedule", "Today's schedule timeline displayed",
     "Has appointments today", "1. View dashboard → schedule section",
     "Timeline shows today's appointments in chronological order", "P2"),
    ("Dashboard", "Quick Actions", "Quick action: Consultation finds first eligible appointment",
     "Has checked_in/confirmed appointment", "1. Tap 'Consultation' quick action",
     "Navigates to consultation for first eligible appointment", "P2"),
    ("Dashboard", "Real-time", "Notification badge updates in real-time",
     "Dashboard open", "1. Trigger notification",
     "Bell icon badge count increments", "P2"),

    # ── Schedule ──
    ("SCHEDULE",),
    ("Schedule", "Calendar", "Calendar view shows appointment dots for dates",
     "Has appointments this month", "1. Go to Schedule tab",
     "Calendar marks dates with appointments using dots/indicators", "P1"),
    ("Schedule", "Calendar", "Navigate between months in calendar",
     "On Schedule screen", "1. Tap forward arrow\n2. Tap back arrow",
     "Month changes, appointments load for new month", "P2"),
    ("Schedule", "Day View", "Tap date shows day's appointments",
     "Has appointments on selected date", "1. Tap a date with appointments",
     "Day appointment list shows all appointments sorted by time with patient info, type, status", "P1"),
    ("Schedule", "Navigation", "Tap appointment navigates to detail",
     "Day appointments shown", "1. Tap an appointment card",
     "Navigates to AppointmentDetail screen", "P1"),

    # ── Appointment Management ──
    ("APPOINTMENT MANAGEMENT",),
    ("Appointment", "Detail", "View appointment detail with patient info and schedule",
     "Has an appointment", "1. Open appointment from schedule",
     "Shows status hero, patient name/ID, date/time/duration/branch, reason/notes", "P1"),
    ("Appointment", "Confirm", "Confirm a requested appointment",
     "Appointment in 'requested' status", "1. Open detail\n2. Tap 'Confirm Appointment'",
     "Status → confirmed, success alert shown", "P1"),
    ("Appointment", "Check In", "Check in a confirmed appointment",
     "Appointment in 'confirmed' status", "1. Open detail\n2. Tap 'Check In Patient'",
     "Status → checked_in, success alert shown", "P1"),
    ("Appointment", "No Show", "Mark appointment as no-show",
     "Appointment confirmed/checked_in", "1. Tap 'Mark No-Show'\n2. Confirm in alert",
     "Status → no_show, confirmation shown", "P1"),
    ("Appointment", "Cancel", "Cancel appointment with reason",
     "Appointment not in terminal status", "1. Tap 'Cancel'\n2. Enter reason in modal\n3. Confirm",
     "Status → cancelled, reason saved, modal closes", "P1"),
    ("Appointment", "Cancel", "Cancel requires reason (validation)",
     "Cancel modal open", "1. Tap Confirm without entering reason",
     "Validation: 'Please provide a reason for cancellation'", "P2"),
    ("Appointment", "Video Call", "Join Video Call button shown for teleconsult",
     "Teleconsult appointment in confirmed/checked_in/in_progress", "1. Open teleconsult appointment detail",
     "'Join Video Call' button visible with video icon", "P1"),
    ("Appointment", "Video Call", "Doctor joins video call and connects to patient",
     "Patient already in video call", "1. Tap 'Join Video Call'\n2. Grant permissions",
     "Video call connects, both see/hear each other", "P1"),
    ("Appointment", "Actions Hidden", "Action buttons hidden for terminal statuses",
     "Completed/cancelled/no_show appointment", "1. Open terminal appointment detail",
     "No action buttons shown, only read-only info + 'View Patient Record' link", "P2"),

    # ── Consultation ──
    ("CONSULTATION",),
    ("Consultation", "Start", "Start consultation from checked_in appointment",
     "Appointment in checked_in/in_progress", "1. Tap 'Start Consultation'",
     "Consultation record created, consultation form opens", "P1"),
    ("Consultation", "Record", "Enter diagnosis, treatment, and notes",
     "In consultation screen", "1. Fill diagnosis field\n2. Fill treatment field\n3. Add notes\n4. Save",
     "Consultation updated, data persisted", "P1"),
    ("Consultation", "Follow-up", "Schedule follow-up from consultation",
     "In active consultation", "1. Tap 'Schedule Follow-Up'\n2. Pick date/time\n3. Enter reason\n4. Confirm",
     "Follow-up appointment created with type='follow_up'", "P1"),
    ("Consultation", "Prescription", "Create prescription from consultation",
     "In active consultation", "1. Tap 'Create Prescription'\n2. Fill medications\n3. Submit",
     "Navigates to PrescriptionForm, prescription created and linked to appointment", "P1"),
    ("Consultation", "Complete", "Complete consultation",
     "In active consultation with data filled", "1. Tap 'Complete Consultation'\n2. Confirm",
     "Consultation completed, appointment status → completed", "P1"),

    # ── Prescriptions ──
    ("PRESCRIPTIONS",),
    ("Prescription", "Create", "Create prescription with multiple medications",
     "In PrescriptionForm", "1. Enter diagnosis\n2. Add medication (name, dosage, frequency, duration)\n3. Add another medication\n4. Submit",
     "Prescription created with all medications, status='active'", "P1"),
    ("Prescription", "Validate", "Prescription requires at least one medication",
     "Empty PrescriptionForm", "1. Try to submit without medications",
     "Validation error: at least one medication required", "P2"),
    ("Prescription", "List", "View issued prescriptions",
     "Has prescriptions", "1. Navigate to PrescriptionList",
     "All doctor's prescriptions shown with patient, medication count, status, date", "P1"),
    ("Prescription", "Detail", "View prescription detail",
     "Has a prescription", "1. Tap a prescription",
     "Full detail with all medications, patient info, diagnosis, status", "P1"),

    # ── Patient Management ──
    ("PATIENT MANAGEMENT",),
    ("Patients", "List", "View patient list with search",
     "Logged in as doctor", "1. Go to Patients tab\n2. Search by name",
     "Patient list with avatar, name, ID, phone; search filters in real-time", "P1"),
    ("Patients", "Record", "View patient record with history",
     "Has patients", "1. Tap a patient",
     "Shows demographics, medical history, prescriptions, appointments, consultations", "P1"),
    ("Patients", "Pagination", "Patient list paginates correctly",
     "More than 20 patients", "1. Scroll to bottom",
     "Next page loads automatically or via button", "P2"),

    # ── Video Call (Doctor Side) ──
    ("VIDEO CALL (DOCTOR)",),
    ("Video Call", "Join", "Doctor joins teleconsult from appointment detail",
     "Teleconsult appointment confirmed", "1. Open appointment\n2. Tap 'Join Video Call'\n3. Grant permissions",
     "Video call screen opens, local camera shown, waiting for patient or connects immediately", "P1"),
    ("Video Call", "Controls", "All video controls work (mute, camera, flip, speaker, end)",
     "In active video call", "1. Test each control button",
     "Mute/unmute toggles audio, camera on/off, flip switches, speaker toggles, end terminates", "P1"),
    ("Video Call", "End", "Ending call updates session status via API",
     "In active call", "1. Tap end call",
     "telemedicineApi.end() called, session status → ended in DB", "P1"),
    ("Video Call", "Remote Left", "Remote participant leaving shows 'Call Ended'",
     "In active call", "1. Patient ends call",
     "Doctor sees 'Call Ended' with duration, auto-navigates back after 2s", "P2"),

    # ── Emergency ──
    ("EMERGENCY",),
    ("Emergency", "View", "View pending emergencies on dashboard",
     "Emergency session exists with status=requested", "1. View dashboard",
     "Emergency alert card shows patient name, type, reason", "P1"),
    ("Emergency", "Accept", "Accept emergency session",
     "Pending emergency", "1. Tap Accept",
     "Session accepted, chat/video opens", "P1"),
    ("Emergency", "Chat", "Respond to emergency via chat",
     "Emergency chat open", "1. Type and send messages",
     "Messages sent and received in real-time", "P1"),
    ("Emergency", "End", "End emergency session",
     "In active emergency", "1. End session",
     "Session status → ended", "P2"),

    # ── Follow-ups ──
    ("FOLLOW-UPS",),
    ("Follow-ups", "List", "View all follow-up appointments",
     "Has follow-up appointments", "1. Navigate to FollowUpList",
     "Shows follow-up appointments sorted by date with patient info", "P1"),
    ("Follow-ups", "Navigate", "Tap follow-up navigates to appointment detail",
     "Has follow-ups", "1. Tap a follow-up",
     "Opens AppointmentDetail for that follow-up", "P2"),

    # ── Profile ──
    ("PROFILE & SETTINGS",),
    ("Profile", "View", "View doctor profile with professional info",
     "Logged in as doctor", "1. Go to Profile tab",
     "Shows avatar, name (Dr. prefix), specialization, department, contact info", "P1"),
    ("Profile", "Edit", "Edit doctor profile",
     "On profile screen", "1. Tap edit\n2. Update specialization/phone\n3. Save",
     "Profile updated", "P2"),
    ("Profile", "Notifications", "Navigate to notification list from profile menu",
     "On profile screen", "1. Tap 'Notifications' menu item",
     "Navigates to NotificationList", "P2"),
    ("Settings", "Theme", "Change app theme",
     "On settings screen", "1. Select Light/Dark/System",
     "Theme changes immediately", "P3"),
]


CLINIC_ADMIN_TESTS = [
    # ── Auth ──
    ("AUTH & ACCESS",),
    ("Auth", "Login", "Clinic admin (super_admin) login",
     "super_admin account in tenant", "1. Login via staff login with super_admin credentials",
     "Logged in, full admin features accessible", "P1"),
    ("Auth", "Role Access", "Verify super_admin can access all clinic-scoped routes",
     "Logged in as super_admin", "1. Access patients, appointments, billing, settings endpoints",
     "All tenant-scoped routes accessible with full permissions", "P1"),

    # ── Branch Management ──
    ("BRANCH MANAGEMENT",),
    ("Branch", "List", "GET /branches — list all branches",
     "Authenticated admin", "1. Call GET /branches",
     "Returns all branches with name, code, address, phone, operating hours, isActive", "P1"),
    ("Branch", "Create", "POST /branches — create new branch",
     "Admin with BRANCH_CREATE", "1. POST /branches with name, code, address, phone, operatingHours",
     "Branch created, returned with _id", "P1"),
    ("Branch", "Update", "PUT /branches/:id — update branch details",
     "Existing branch", "1. PUT with updated name, phone, hours",
     "Branch updated", "P2"),
    ("Branch", "Delete", "DELETE /branches/:id — deactivate branch",
     "Existing branch", "1. DELETE /branches/:id",
     "Branch soft-deleted (isActive=false or deletedAt set)", "P2"),

    # ── Department Management ──
    ("DEPARTMENT MANAGEMENT",),
    ("Department", "List", "GET /departments — list all departments",
     "Authenticated admin", "1. Call GET /departments",
     "Returns departments with name, code, branch, head of department", "P1"),
    ("Department", "Create", "POST /departments — create department",
     "Admin with DEPARTMENT_CREATE", "1. POST with name, code, branchId, headOfDepartment",
     "Department created", "P1"),
    ("Department", "Update", "PUT /departments/:id — update department",
     "Existing department", "1. PUT with updated fields",
     "Department updated", "P2"),
    ("Department", "Delete", "DELETE /departments/:id — remove department",
     "Existing department", "1. DELETE",
     "Department removed", "P2"),

    # ── Doctor/Staff User Management ──
    ("USER MANAGEMENT",),
    ("User", "List Doctors", "GET /doctors — list doctors",
     "Admin with USER_READ", "1. GET /doctors",
     "Returns all doctor users with name, specialization, department, branch", "P1"),
    ("User", "Create Doctor", "POST /doctors — create doctor account",
     "Admin with USER_CREATE", "1. POST with name, email, phone, specialization, department, branch",
     "Doctor user created, can login", "P1"),
    ("User", "Update Doctor", "PUT /doctors/:id — update doctor",
     "Existing doctor", "1. PUT with updated fields",
     "Doctor updated", "P2"),
    ("User", "Delete Doctor", "DELETE /doctors/:id — deactivate doctor",
     "Existing doctor", "1. DELETE",
     "Doctor deactivated", "P2"),
    ("User", "List Staff", "GET /staff — list staff members",
     "Admin with USER_READ", "1. GET /staff",
     "Returns all staff with roles", "P1"),
    ("User", "Create Staff", "POST /staff — create staff member",
     "Admin with USER_CREATE", "1. POST with name, email, role (receptionist, billing_officer, etc.)",
     "Staff created with correct role and permissions", "P1"),
    ("User", "Roles", "GET /roles — view available roles and permissions",
     "Admin", "1. GET /roles",
     "Returns all roles with permission sets", "P2"),

    # ── Patient Management ──
    ("PATIENT MANAGEMENT",),
    ("Patient", "List", "GET /patients — list all patients",
     "Admin with PATIENT_READ", "1. GET /patients with pagination",
     "Paginated patient list with search", "P1"),
    ("Patient", "Create", "POST /patients — create patient",
     "Admin with PATIENT_CREATE", "1. POST with name, email, phone, DOB, gender",
     "Patient created with auto-generated patientId", "P1"),
    ("Patient", "Update", "PUT /patients/:id — update patient",
     "Existing patient", "1. PUT with updated fields",
     "Patient updated", "P2"),
    ("Patient", "Delete", "DELETE /patients/:id — soft-delete patient",
     "Existing patient", "1. DELETE",
     "Patient soft-deleted (deletedAt set)", "P2"),

    # ── Appointment Management ──
    ("APPOINTMENT MANAGEMENT",),
    ("Appointment", "List", "GET /appointments — list all clinic appointments",
     "Admin with APPOINTMENT_READ", "1. GET /appointments with date range filters",
     "All appointments visible (not just own)", "P1"),
    ("Appointment", "Create", "POST /appointments — create appointment on behalf of patient",
     "Admin with APPOINTMENT_CREATE", "1. POST with patient, doctor, scheduledAt, type",
     "Appointment created", "P1"),
    ("Appointment", "Update", "PUT /appointments/:id — update any appointment",
     "Admin with APPOINTMENT_UPDATE", "1. PUT to change status/notes",
     "Appointment updated", "P1"),

    # ── Billing & Invoices ──
    ("BILLING & INVOICES",),
    ("Invoice", "List", "GET /invoices — list all invoices",
     "Admin with INVOICE_READ", "1. GET /invoices with status filter",
     "All invoices visible with amount, status, patient, dates", "P1"),
    ("Invoice", "Create", "POST /invoices — create invoice",
     "Admin with INVOICE_CREATE", "1. POST with patient, items[], amounts",
     "Invoice created with auto-generated number", "P1"),
    ("Invoice", "Update", "PUT /invoices/:id — update invoice",
     "Draft/sent invoice", "1. PUT to update items or amounts",
     "Invoice updated", "P2"),
    ("Invoice", "PDF", "GET /invoices/:id/pdf — download PDF",
     "Existing invoice", "1. GET PDF endpoint",
     "PDF generated and returned", "P2"),
    ("Billing", "Service Items", "CRUD service items for billing",
     "Admin with SERVICE_ITEM_*", "1. GET/POST/PUT service items",
     "Service items manageable (description, price, tax)", "P2"),
    ("Billing", "Payment Record", "POST /billing/invoice/:id/payment — record payment",
     "Invoice exists", "1. POST payment with amount, method",
     "Payment recorded, invoice status updated (paid if full)", "P1"),
    ("Billing", "Analytics", "GET /billing/analytics/revenue — revenue report",
     "Admin with REVENUE_REPORT_VIEW", "1. GET revenue analytics with date range",
     "Revenue breakdown by period, doctor, department", "P2"),
    ("Billing", "Ledger", "GET /billing/ledger/:patientId — patient ledger",
     "Admin with LEDGER_READ", "1. GET ledger",
     "Full financial history for patient (invoices, payments, advances, adjustments)", "P2"),
    ("Billing", "Advance", "POST /billing/ledger/advance — record advance payment",
     "Admin with PAYMENT_CREATE", "1. POST advance amount",
     "Advance recorded in ledger", "P2"),

    # ── Payment Gateway ──
    ("PAYMENT GATEWAY",),
    ("Payment", "Create Order", "POST /payments/gateway/order — create Razorpay order",
     "Admin with PAYMENT_CREATE", "1. POST with invoiceId, amount",
     "Razorpay order created, order_id returned", "P1"),
    ("Payment", "Verify", "POST /payments/gateway/verify — verify payment",
     "Order created and paid", "1. POST with razorpay_payment_id, order_id, signature",
     "Payment verified, invoice marked paid", "P1"),
    ("Payment", "Webhook", "POST /payments/gateway/webhook — Razorpay webhook",
     "Payment event triggered", "1. Razorpay sends webhook",
     "Payment event processed, signature verified, status updated", "P1"),
    ("Payment", "Refund", "POST /payments/:id/refund — refund payment",
     "Paid payment exists", "1. POST refund",
     "Refund processed, invoice status updated", "P2"),

    # ── Cash Register ──
    ("CASH REGISTER",),
    ("Cash Register", "Open", "POST /cash-register/open — open register",
     "Admin with CASH_REGISTER_OPEN", "1. POST with opening balance",
     "Register opened with timestamp", "P2"),
    ("Cash Register", "Close", "PUT /cash-register/:id/close — close register",
     "Register open", "1. PUT close with closing balance",
     "Register closed, transactions finalized", "P2"),
    ("Cash Register", "Reconcile", "PUT /cash-register/:id/reconcile",
     "Closed register", "1. PUT reconcile with actual amounts",
     "Discrepancy calculated, register reconciled", "P3"),

    # ── Credit Notes ──
    ("CREDIT NOTES",),
    ("Credit Note", "Create", "POST /credit-notes — create credit note",
     "Admin with CREDIT_NOTE_CREATE", "1. POST with patient, amount, reason",
     "Credit note created in draft status", "P2"),
    ("Credit Note", "Issue", "PUT /credit-notes/:id/issue — approve credit note",
     "Draft credit note", "1. PUT issue",
     "Credit note issued, available for use", "P2"),
    ("Credit Note", "Apply", "PUT /credit-notes/:id/apply — apply to invoice",
     "Issued credit note + pending invoice", "1. PUT apply with invoiceId",
     "Credit applied, invoice balance reduced", "P2"),

    # ── Consent Management ──
    ("CONSENT MANAGEMENT",),
    ("Consent", "List", "GET /consents — list all consents",
     "Admin with CONSENT_READ", "1. GET /consents",
     "All consent records visible", "P2"),
    ("Consent", "Create", "POST /consents — create consent request",
     "Admin with CONSENT_CREATE", "1. POST with patient, type, scope, purpose",
     "Consent request created, patient can approve", "P2"),

    # ── Insurance ──
    ("INSURANCE",),
    ("Insurance", "Providers", "CRUD insurance providers",
     "Admin with INSURANCE_*", "1. GET/POST/PUT/DELETE providers",
     "Providers manageable", "P2"),
    ("Insurance", "Claims", "Create and manage insurance claims",
     "Admin with INSURANCE_CREATE", "1. POST claim\n2. PUT status\n3. Reconcile",
     "Claim lifecycle works: create → pre-auth → approve → settle → reconcile", "P2"),
    ("Insurance", "Co-pay", "GET /insurance/claims/:id/co-pay — calculate co-pay",
     "Claim with pre-auth", "1. GET co-pay calculation",
     "Returns patient co-pay amount after insurance coverage", "P3"),

    # ── Queue Management ──
    ("QUEUE MANAGEMENT",),
    ("Queue", "Add", "POST /queue — add patient to queue",
     "Admin with QUEUE_MANAGE", "1. POST with patient, appointment, branch",
     "Patient added to queue", "P2"),
    ("Queue", "Call", "PUT /queue/:id/call — call next patient",
     "Patient in queue", "1. PUT call",
     "Patient called, queue updated", "P2"),
    ("Queue", "Complete", "PUT /queue/:id/complete — complete queue entry",
     "Called patient", "1. PUT complete",
     "Queue entry completed", "P2"),

    # ── Analytics & Reports ──
    ("ANALYTICS & REPORTS",),
    ("Analytics", "Summary", "GET /analytics/summary — clinic summary",
     "Admin with CLINIC_ANALYTICS_VIEW", "1. GET summary",
     "Returns high-level KPIs", "P1"),
    ("Analytics", "Revenue", "GET /analytics/revenue — revenue trends (requires advancedReports)",
     "Feature guard: advancedReports=true", "1. GET /analytics/revenue with date range",
     "Revenue trend data returned; blocked if feature disabled", "P2"),
    ("Analytics", "Appointments", "GET /analytics/appointments — appointment analytics",
     "advancedReports enabled", "1. GET analytics",
     "Appointment volume, type distribution, no-show rates", "P2"),
    ("Analytics", "Feature Guard", "Analytics blocked when advancedReports disabled",
     "Tenant without advancedReports", "1. GET /analytics/revenue",
     "403: Feature not available on current plan", "P2"),

    # ── Subscription Info ──
    ("SUBSCRIPTION",),
    ("Subscription", "View", "GET /subscriptions — view current subscription",
     "Admin with SUBSCRIPTION_READ", "1. GET /subscriptions",
     "Returns current plan, status, trial dates, features, limits", "P1"),
    ("Subscription", "Feature Limits", "Verify feature limits enforce correctly",
     "Subscription with limits", "1. Try to create doctor beyond maxDoctors limit",
     "Error: Plan limit reached for doctors", "P1"),

    # ── Notifications ──
    ("NOTIFICATIONS",),
    ("Notification", "List", "GET /notifications — list notifications",
     "Admin", "1. GET /notifications",
     "Returns notification list with filtering", "P1"),
    ("Notification", "Mark Read", "PUT /notifications/read-all",
     "Has unread notifications", "1. PUT read-all",
     "All marked as read", "P2"),

    # ── Payroll (if accessible) ──
    ("PAYROLL",),
    ("Payroll", "Salary Structures", "CRUD salary structures",
     "Admin with PAYROLL_MANAGE", "1. GET/POST/PUT salary structures",
     "Salary structures manageable", "P3"),
    ("Payroll", "Generate Payslips", "POST /payroll/payslips/generate",
     "Admin with PAYROLL_MANAGE", "1. POST generate with month/year",
     "Payslips generated for all staff", "P3"),
    ("Payroll", "Approve & Pay", "PUT /payroll/payslips/:id/approve then /pay",
     "Generated payslip", "1. Approve\n2. Mark as paid",
     "Payslip approved then marked paid", "P3"),

    # ── Integrations ──
    ("INTEGRATIONS",),
    ("Integration", "List", "GET /integrations — list integrations (requires apiAccess)",
     "Feature guard: apiAccess=true", "1. GET /integrations",
     "Returns configured integrations; blocked if apiAccess disabled", "P3"),
    ("Integration", "Configure", "PUT /integrations — configure integration",
     "Admin with integration:manage + apiAccess", "1. PUT with type, config",
     "Integration saved", "P3"),
    ("Integration", "Test", "POST /integrations/:type/test — test connection",
     "Configured integration", "1. POST test",
     "Connection test result returned", "P3"),
]


PLATFORM_ADMIN_TESTS = [
    # ── Auth ──
    ("AUTH",),
    ("Auth", "Login", "POST /auth/platform/login — platform admin login",
     "Platform user exists", "1. POST with email, password (no tenant header)",
     "JWT returned with role=product_owner, signed with JWT_PLATFORM_SECRET", "P1"),
    ("Auth", "Rejection", "Tenant user cannot access platform routes",
     "Logged in as tenant admin", "1. Try GET /platform/dashboard",
     "403 Forbidden", "P1"),
    ("Auth", "Create User", "POST /platform/users — create platform admin",
     "Logged in as platform admin", "1. POST with name, email, password",
     "New platform user created", "P2"),

    # ── Dashboard ──
    ("DASHBOARD",),
    ("Dashboard", "KPIs", "GET /platform/dashboard — platform KPIs",
     "Authenticated platform admin", "1. GET /platform/dashboard",
     "Returns totalTenants, activeTenants, trialTenants, suspendedTenants, MRR, plan distribution, recent signups", "P1"),

    # ── Tenant Lifecycle ──
    ("TENANT MANAGEMENT",),
    ("Tenant", "Create", "POST /platform/tenants — create new clinic tenant",
     "Platform admin", "1. POST with name, slug, email, phone, address, superAdmin info",
     "Tenant created in PENDING status, provisioning job queued", "P1"),
    ("Tenant", "Create", "Create tenant with duplicate slug fails",
     "Slug already taken", "1. POST with existing slug",
     "Error: slug already in use", "P1"),
    ("Tenant", "List", "GET /platform/tenants — list tenants with filters",
     "Platform admin", "1. GET with ?status=active&plan=starter&search=clinic",
     "Paginated list filtered by status, plan, search", "P1"),
    ("Tenant", "Detail", "GET /platform/tenants/:slug — get tenant details",
     "Existing tenant", "1. GET by slug",
     "Full tenant details returned", "P1"),
    ("Tenant", "Update", "PUT /platform/tenants/:slug — update tenant metadata",
     "Existing tenant", "1. PUT with updated name, email, phone",
     "Tenant updated, slug immutable", "P1"),
    ("Tenant", "Activate", "POST /platform/tenants/:slug/activate",
     "Tenant in PENDING status with superAdmin assigned", "1. POST activate",
     "Tenant activated, isActive=true, Redis cache invalidated", "P1"),
    ("Tenant", "Activate", "Activate fails if no superAdmin assigned",
     "Tenant without superAdmin", "1. POST activate",
     "Error: SuperAdmin must be assigned first", "P2"),
    ("Tenant", "Suspend", "POST /platform/tenants/:slug/suspend",
     "Active tenant", "1. POST with { reason }",
     "Tenant suspended, isActive=false, all API access blocked for tenant users", "P1"),
    ("Tenant", "Suspend Effect", "Verify suspended tenant users are blocked",
     "Suspended tenant", "1. Login as tenant user\n2. Try any API call",
     "Login works (to see suspension message), all other routes return 403 with SUBSCRIPTION_BLOCKED", "P1"),
    ("Tenant", "Delete", "DELETE /platform/tenants/:slug — soft delete",
     "Existing tenant", "1. DELETE without hardDelete",
     "Tenant soft-deleted (deletedAt set), hidden from queries", "P1"),
    ("Tenant", "Hard Delete", "DELETE /platform/tenants/:slug — hard delete",
     "Soft-deleted tenant", "1. DELETE with { hardDelete: true }",
     "Tenant permanently removed (irreversible)", "P2"),
    ("Tenant", "Features", "PUT /platform/tenants/:slug/features — update features",
     "Active tenant", "1. PUT with { features: { maxDoctors: 10, teleHealthEnabled: true, ... } }",
     "Tenant features updated, limits enforced on next check", "P1"),
    ("Tenant", "SuperAdmin", "POST /platform/tenants/:slug/superadmin — assign admin",
     "Tenant without or with existing superAdmin", "1. POST with { name, email, phone }",
     "SuperAdmin assigned/reassigned, audit logged", "P1"),
    ("Tenant", "Usage", "GET /platform/tenants/:slug/usage",
     "Active tenant", "1. GET usage",
     "Returns doctors, staff, branches, patients, storage, SMS counts", "P2"),
    ("Tenant", "Audit Log", "GET /platform/tenants/:slug/audit-log",
     "Tenant with history", "1. GET audit log with pagination",
     "Returns all audit entries with action, performedBy, before/after, reason, timestamp", "P1"),

    # ── Provisioning ──
    ("PROVISIONING",),
    ("Provisioning", "Status", "GET /platform/tenants/:slug/provisioning",
     "Newly created tenant", "1. GET provisioning status",
     "Returns status (pending/in_progress/completed/failed), attempt count, lastError", "P1"),
    ("Provisioning", "Retry", "POST /platform/tenants/:slug/provisioning/retry",
     "Failed provisioning", "1. POST retry",
     "New provisioning job queued, status reset", "P1"),
    ("Provisioning", "Rollback", "POST /platform/tenants/:slug/provisioning/rollback",
     "Failed provisioning", "1. POST rollback",
     "Rollback job queued, partial DB cleanup", "P2"),
    ("Provisioning", "Complete", "Verify full provisioning flow",
     "New tenant created", "1. Create tenant\n2. Poll provisioning status\n3. Wait for completion",
     "Status: pending → in_progress → completed; tenant DB created with collections + indexes", "P1"),

    # ── Subscription Management ──
    ("SUBSCRIPTION MANAGEMENT",),
    ("Subscription", "Create", "POST /platform/tenants/:slug/subscription",
     "Tenant without subscription", "1. POST with planId, billingCycle, trialDays",
     "Subscription created with trial dates, status=trial", "P1"),
    ("Subscription", "Change Plan", "PUT /platform/tenants/:slug/subscription/plan",
     "Active subscription", "1. PUT with newPlanId, billingCycle",
     "Plan changed, proration calculated, status updated", "P1"),
    ("Subscription", "Cancel", "POST /platform/tenants/:slug/subscription/cancel",
     "Active subscription", "1. POST with { reason }",
     "Subscription cancelled, tenant access blocked after grace", "P1"),
    ("Subscription", "Reactivate", "POST /platform/tenants/:slug/subscription/reactivate",
     "Cancelled subscription", "1. POST reactivate",
     "Subscription reactivated, tenant access restored", "P1"),
    ("Subscription", "Custom Pricing", "PUT /platform/tenants/:slug/subscription/custom-pricing",
     "Active subscription", "1. PUT with basePrice, discountPercent, reason",
     "Custom pricing set, audit logged", "P2"),
    ("Subscription", "Extend Trial", "POST /platform/subscriptions/:id/extend-trial",
     "Trial subscription", "1. POST with { additionalDays: 14 }",
     "Trial extended by 14 days, trialEndsAt updated", "P1"),
    ("Subscription", "List", "GET /platform/subscriptions — list all subscriptions",
     "Platform admin", "1. GET with ?status=active",
     "All subscriptions listed with plan, tenant, status, dates", "P1"),
    ("Subscription", "Expiring", "GET /platform/subscriptions/expiring-trials?days=7",
     "Subscriptions nearing expiry", "1. GET expiring",
     "Returns trials expiring within 7 days", "P2"),
    ("Subscription", "Cancel Change", "DELETE /platform/tenants/:slug/subscription/scheduled-change",
     "Scheduled plan change pending", "1. DELETE",
     "Scheduled change cancelled", "P3"),

    # ── Plan Management ──
    ("PLAN MANAGEMENT",),
    ("Plan", "List", "GET /platform/plans — list all plans",
     "Platform admin", "1. GET /platform/plans",
     "Returns all plans with pricing, features, limits", "P1"),
    ("Plan", "Create", "POST /platform/plans — create pricing plan",
     "Platform admin", "1. POST with name, slug, tier, priceMonthly, priceYearly, features",
     "Plan created; slug must be unique", "P1"),
    ("Plan", "Update", "PUT /platform/plans/:id — update plan",
     "Existing plan", "1. PUT with updated pricing/features",
     "Plan updated (slug immutable)", "P1"),
    ("Plan", "Delete", "DELETE /platform/plans/:id — deactivate plan",
     "Existing plan", "1. DELETE",
     "Plan deactivated (isActive=false), existing subscribers unaffected", "P2"),
    ("Plan", "Custom", "POST /platform/tenants/:slug/custom-plan",
     "Specific tenant", "1. POST with custom features/pricing",
     "Custom plan created (slug=custom-{clinicSlug}-{timestamp})", "P2"),
    ("Plan", "Compare", "POST /platform/plans/compare — side-by-side comparison",
     "Two plans exist", "1. POST with planIdA, planIdB",
     "Returns tierDirection, pricing diff, feature diff, limit diff", "P2"),

    # ── Pricing ──
    ("PRICING",),
    ("Pricing", "Calculate", "POST /platform/pricing/calculate",
     "Plan exists", "1. POST with planId, billingCycle, extraDoctors, promoCode",
     "Returns basePrice, addOns, discount, tax, total", "P1"),
    ("Pricing", "Proration", "POST /platform/pricing/proration/:subscriptionId",
     "Active subscription", "1. POST with newPlanId",
     "Returns currentValue, newValue, credit/charge, daysRemaining", "P2"),

    # ── Promotions ──
    ("PROMOTIONS",),
    ("Promotion", "Create", "POST /platform/promotions — create promo code",
     "Platform admin", "1. POST with code, discountType, discountValue, validFrom/Until, maxUses",
     "Promotion created with unique code", "P1"),
    ("Promotion", "Validate", "POST /platform/promotions/validate",
     "Valid promo code", "1. POST with { code, planId, billingCycle, clinicSlug }",
     "Returns valid=true, discount details, remaining uses", "P1"),
    ("Promotion", "Validate Expired", "Validate expired promo code",
     "Expired promo", "1. POST validate",
     "Returns valid=false with reason", "P2"),
    ("Promotion", "Analytics", "GET /platform/promotions/analytics",
     "Promotions with usage", "1. GET analytics",
     "Redemption count, revenue impact, conversion rate", "P2"),
    ("Promotion", "Update", "PUT /platform/promotions/:id",
     "Existing promotion", "1. PUT with updated fields",
     "Promotion updated", "P2"),
    ("Promotion", "Delete", "DELETE /platform/promotions/:id",
     "Existing promotion", "1. DELETE",
     "Promotion deactivated", "P3"),

    # ── Rebates ──
    ("REBATES",),
    ("Rebate", "CRUD", "Full CRUD for rebate rules",
     "Platform admin", "1. POST create\n2. GET list\n3. PUT update\n4. DELETE",
     "Rebate lifecycle works: volume/loyalty/partner types", "P3"),

    # ── Referrals ──
    ("REFERRALS",),
    ("Referral", "Generate", "POST /platform/tenants/:slug/referral-code",
     "Active tenant", "1. POST generate",
     "Unique referral code generated with expiry and discount", "P2"),
    ("Referral", "Lookup", "POST /platform/referrals/lookup",
     "Valid referral code", "1. POST with { referralCode }",
     "Returns code details, issuing tenant, discount", "P2"),
    ("Referral", "Apply", "POST /platform/tenants/:slug/apply-referral",
     "Valid referral code + new tenant", "1. POST with { referralCode }",
     "Discount applied to tenant's subscription", "P2"),

    # ── Feature Flags ──
    ("FEATURE FLAGS",),
    ("Feature Flag", "List", "GET /platform/feature-flags",
     "Platform admin", "1. GET all flags",
     "Returns global flags with keys, values, tenant overrides", "P1"),
    ("Feature Flag", "Create", "POST /platform/feature-flags",
     "Platform admin", "1. POST with key, name, category, enabled",
     "Flag created (key must be unique)", "P1"),
    ("Feature Flag", "Toggle", "PUT /platform/feature-flags/:id — enable/disable",
     "Existing flag", "1. PUT with { enabled: false }",
     "Flag disabled globally", "P1"),
    ("Feature Flag", "Override", "POST /platform/feature-flags/tenant-override",
     "Flag + tenant slug", "1. POST with { flagId, clinicSlug, value: true }",
     "Tenant override set, takes precedence over global", "P1"),
    ("Feature Flag", "Resolve", "GET /platform/feature-flags/resolve/:slug",
     "Tenant with flags + overrides", "1. GET resolve for tenant",
     "Returns merged flag values (global + overrides)", "P1"),
    ("Feature Flag", "Remove Override", "POST /platform/feature-flags/tenant-override/remove",
     "Tenant with override", "1. POST remove",
     "Override removed, falls back to global value", "P2"),

    # ── Analytics ──
    ("ANALYTICS",),
    ("Analytics", "Revenue", "GET /platform/analytics/revenue",
     "Platform admin", "1. GET with startDate, endDate",
     "Monthly revenue trend, churned count, new subscriptions", "P1"),
    ("Analytics", "MRR", "GET /platform/analytics/mrr",
     "Platform admin", "1. GET MRR breakdown",
     "MRR by plan tier", "P1"),
    ("Analytics", "Subscriptions", "GET /platform/analytics/subscriptions",
     "Platform admin", "1. GET subscription status breakdown",
     "Active/trial/grace/suspended/cancelled counts", "P1"),
    ("Analytics", "Expiring", "GET /platform/analytics/expiring?days=30",
     "Subscriptions exist", "1. GET expiring",
     "Subscriptions expiring in 30 days listed", "P2"),
    ("Analytics", "Suspended", "GET /platform/analytics/suspended",
     "Suspended tenants exist", "1. GET suspended",
     "Suspended tenants with reason and date", "P2"),
    ("Analytics", "Plan Distribution", "GET /platform/analytics/plans",
     "Multiple tenants on different plans", "1. GET plans",
     "Count per plan tier", "P2"),
    ("Analytics", "Tenant Growth", "GET /platform/analytics/tenant-growth",
     "Platform admin", "1. GET with period=monthly",
     "Monthly new tenant signup trend", "P2"),
    ("Analytics", "Churn", "GET /platform/analytics/churn",
     "Cancelled subscriptions exist", "1. GET with date range",
     "Churn rate, reasons breakdown", "P2"),
    ("Analytics", "Feature Adoption", "GET /platform/analytics/feature-adoption",
     "Multiple tenants", "1. GET feature adoption",
     "Feature usage across tenants (teleHealth, advancedReports, etc.)", "P3"),
    ("Analytics", "Promotion Impact", "GET /platform/analytics/promotion-impact",
     "Promotions with usage", "1. GET with promoId",
     "Revenue impact of specific promotion", "P3"),

    # ── Monitoring ──
    ("MONITORING",),
    ("Monitoring", "Health", "GET /monitoring/health (public)",
     "None", "1. GET /monitoring/health",
     "Returns { status: 'ok' }", "P1"),
    ("Monitoring", "Metrics", "GET /monitoring/metrics (platform auth)",
     "Platform admin", "1. GET /monitoring/metrics",
     "API request counts, latencies, error rates", "P2"),
    ("Monitoring", "System", "GET /monitoring/system",
     "Platform admin", "1. GET system report",
     "CPU, memory, uptime, node version", "P2"),
    ("Monitoring", "Errors", "GET /monitoring/errors",
     "Platform admin", "1. GET recent errors",
     "Recent error log entries", "P2"),
    ("Monitoring", "Slow Queries", "GET /monitoring/slow-queries",
     "Platform admin", "1. GET slow queries",
     "DB queries exceeding threshold", "P3"),

    # ── Public Registration Flow ──
    ("PUBLIC REGISTRATION (Self-Service)",),
    ("Registration", "List Plans", "GET /public/plans",
     "None (public)", "1. GET /public/plans",
     "Returns all public, active plans with pricing", "P1"),
    ("Registration", "Check Slug", "GET /public/check-slug/:slug",
     "None", "1. GET with desired slug",
     "Returns { available: true/false }", "P1"),
    ("Registration", "Create Order", "POST /public/create-registration-order",
     "Paid plan selected", "1. POST with planSlug, billingCycle",
     "Razorpay order created, orderId returned", "P1"),
    ("Registration", "Register", "POST /public/register-clinic",
     "Order paid or free plan", "1. POST with clinic details, admin info, orderId",
     "Clinic registered, provisioning started, tracking token returned", "P1"),
    ("Registration", "Track", "GET /public/provisioning-status/:token",
     "Provisioning in progress", "1. Poll status",
     "Returns step-by-step progress until completion", "P1"),
    ("Registration", "Free Plan", "Register clinic on free plan (no payment)",
     "Free plan exists", "1. POST register-clinic without orderId\n2. Track provisioning",
     "Clinic created on free plan, no payment required", "P1"),
]


# ═══════════════════════════════════════════════════════════
#  GENERATE WORKBOOK
# ═══════════════════════════════════════════════════════════

def main():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    sheets = [
        ("Patient", "PAT", PATIENT_TESTS),
        ("Doctor & Staff", "DOC", DOCTOR_TESTS),
        ("Clinic Admin", "ADM", CLINIC_ADMIN_TESTS),
        ("Platform Admin", "PLT", PLATFORM_ADMIN_TESTS),
    ]

    for sheet_name, prefix, tests in sheets:
        ws = wb.create_sheet(title=sheet_name)
        setup_sheet(ws, sheet_name, COLUMNS)
        add_tests(ws, prefix, tests)

    # ── Summary Sheet ──
    ws_sum = wb.create_sheet(title="Summary", index=0)
    ws_sum.sheet_properties.tabColor = "333333"

    # Title
    ws_sum.merge_cells("A1:F1")
    title_cell = ws_sum["A1"]
    title_cell.value = "BestoClinic — UAT Test Case Summary"
    title_cell.font = Font(name="Calibri", bold=True, size=16, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center")

    ws_sum.merge_cells("A2:F2")
    ws_sum["A2"].value = f"Generated: {date.today().isoformat()}  |  Internal Testing & QA Progress Tracker"
    ws_sum["A2"].font = Font(name="Calibri", size=10, italic=True, color="666666")
    ws_sum["A2"].alignment = Alignment(horizontal="center")

    # Count stats
    summary_headers = ["Sheet", "Role", "Total Test Cases", "P1 (Critical)", "P2 (Important)", "P3 (Nice-to-have)"]
    for col_idx, h in enumerate(summary_headers, 1):
        cell = ws_sum.cell(row=4, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    widths = [20, 22, 18, 16, 16, 18]
    for i, w in enumerate(widths, 1):
        ws_sum.column_dimensions[get_column_letter(i)].width = w

    row = 5
    grand_total = 0
    grand_p1 = 0
    grand_p2 = 0
    grand_p3 = 0

    for sheet_name, prefix, tests in sheets:
        total = sum(1 for t in tests if len(t) > 1)
        p1 = sum(1 for t in tests if len(t) > 1 and t[6] == "P1")
        p2 = sum(1 for t in tests if len(t) > 1 and t[6] == "P2")
        p3 = sum(1 for t in tests if len(t) > 1 and t[6] == "P3")
        grand_total += total
        grand_p1 += p1
        grand_p2 += p2
        grand_p3 += p3

        role_map = {
            "Patient": "Patient / Guardian",
            "Doctor & Staff": "Doctor / Receptionist / Staff",
            "Clinic Admin": "SuperAdmin / Admin / Manager",
            "Platform Admin": "Product Owner",
        }
        data = [sheet_name, role_map[sheet_name], total, p1, p2, p3]
        for col_idx, val in enumerate(data, 1):
            cell = ws_sum.cell(row=row, column=col_idx, value=val)
            cell.font = BODY_FONT
            cell.alignment = CENTER
            cell.border = THIN_BORDER
        row += 1

    # Grand total
    totals = ["", "GRAND TOTAL", grand_total, grand_p1, grand_p2, grand_p3]
    for col_idx, val in enumerate(totals, 1):
        cell = ws_sum.cell(row=row, column=col_idx, value=val)
        cell.font = Font(name="Calibri", bold=True, size=11)
        cell.alignment = CENTER
        cell.border = THIN_BORDER
        cell.fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")

    # Legend
    row += 2
    ws_sum.cell(row=row, column=1, value="Priority Legend:").font = Font(bold=True, size=10)
    row += 1
    for label, fill, desc in [
        ("P1", P1_FILL, "Critical — Must pass for release. Core functionality."),
        ("P2", P2_FILL, "Important — Should pass. Significant but not blocking."),
        ("P3", P3_FILL, "Nice-to-have — Can defer. Edge cases, cosmetic."),
    ]:
        ws_sum.cell(row=row, column=1, value=label).fill = fill
        ws_sum.cell(row=row, column=1).font = Font(bold=True, size=10)
        ws_sum.cell(row=row, column=2, value=desc).font = Font(size=10)
        ws_sum.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        row += 1

    row += 1
    ws_sum.cell(row=row, column=1, value="Status Options:").font = Font(bold=True, size=10)
    row += 1
    for status in ["Not Tested", "Pass", "Fail", "Blocked", "Deferred"]:
        ws_sum.cell(row=row, column=1, value=status).font = Font(size=10)
        row += 1

    # Save
    output_path = "/Users/akashkamble/Downloads/BestoClinic/BestoClinic_UAT_TestCases.xlsx"
    wb.save(output_path)
    print(f"Generated: {output_path}")
    print(f"Total test cases: {grand_total} (P1:{grand_p1} P2:{grand_p2} P3:{grand_p3})")


if __name__ == "__main__":
    main()
